"""
Agent gateway tests — the ``gateway_*`` functions in ``tracebi/mcp_server.py``.

The gateway's promise is the *stamp*: every query response carries the
resolved query, the lineage chain, and a fingerprint of the full result.
These tests hold it to that — a stamped response must re-verify, a preview
cap must not change the fingerprint, and JSON must survive the trip —
because an agent will quote these numbers to a person who never sees this
code.

The MCP registration layer itself is only exercised when the optional
``mcp`` package is installed (skipped otherwise); the operations are plain
functions precisely so the suite does not depend on it.
"""

import json

import pandas as pd
import pytest

from tracebi import DataModel, MemoryConnector
from tracebi import model_registry
from tracebi.mcp_server import (
    gateway_query,
)


@pytest.fixture()
def gateway_model():
    """A small star schema registered under a name no other test uses."""
    orders = pd.DataFrame({
        "order_id":    [1, 2, 3, 4, 5, 6],
        "customer_id": [1, 2, 1, 3, 2, 1],
        "revenue":     [100.0, 250.0, 75.0, 300.0, 125.0, 50.0],
        "status":      ["shipped", "open", "shipped", "shipped", "open", "shipped"],
    })
    customers = pd.DataFrame({
        "customer_id": [1, 2, 3],
        "region":      ["NE", "SE", "MW"],
    })
    connector = MemoryConnector("gw_mem", tables={
        "orders": orders, "customers": customers,
    })
    model = DataModel("gw_demo")
    model.add_connector(connector)
    model.add_table("orders", connector="gw_mem", source="orders")
    model.add_table("customers", connector="gw_mem", source="customers")
    model.add_relationship(
        "orders_to_customers",
        left_table="orders", right_table="customers", left_key="customer_id",
    )
    model.add_dimension(
        "dim_customer", table_name="customers",
        key_col="customer_id", attributes=["region"],
    )
    model.add_fact(
        "fact_orders", table_name="orders",
        measures=["revenue"], foreign_keys={"dim_customer": "customer_id"},
    )
    model.add_measure(
        "total_revenue", column="revenue", agg="sum", format="currency",
    )
    model.connect()
    model_registry.register(model)
    return model


def _query(**overrides):
    kwargs = dict(
        model="gw_demo",
        fact="fact_orders",
        measures={"revenue": "sum"},
        dimensions=["dim_customer.region"],
    )
    kwargs.update(overrides)
    return gateway_query(**kwargs)


# ── The stamp ──────────────────────────────────────────────────────────────


def test_stamp_reverifies(gateway_model):
    """The recorded query, re-run, reproduces the recorded fingerprint."""
    out = _query()
    q = out["query"]
    ds = gateway_model.query(
        fact=q["fact"], measures=q["measures"],
        dimensions=q["dimensions"], filters=q["filters"] or None,
        aggregate=q["aggregate"], allow_fanout=q["allow_fanout"],
    )
    assert ds.fingerprint() == out["fingerprint"]


def test_preview_cap_does_not_change_the_fingerprint(gateway_model):
    """rows is transport; the stamp covers the full result."""
    full = _query()
    capped = _query(preview_rows=1)
    assert capped["rows_returned"] == 1
    assert capped["truncated"]
    assert capped["row_count"] == full["row_count"]
    assert capped["fingerprint"] == full["fingerprint"]


def test_order_by_tiebreak_is_noted_and_the_fingerprint_is_unchanged(gateway_model):
    """Tie-break keys stay in the receipt. The note says the gateway added them."""
    import anyio

    from tracebi.mcp_server import build_server

    out = _query(order_by=["-revenue"])
    cols = [o["column"] for o in out["query"]["order_by"]]
    assert cols[0] == "revenue"
    assert "dim_customer.region" in cols[1:]
    assert out["binding"]["query"]["order_by"] == out["query"]["order_by"]
    assert "tie-break" in out["order_by_note"]
    direct = gateway_model.query(
        fact="fact_orders", measures={"revenue": "sum"},
        dimensions=["dim_customer.region"], order_by=["-revenue"],
    )
    assert direct.fingerprint() == out["fingerprint"]
    plain = _query()
    assert "order_by" not in plain["query"]
    assert "order_by_note" not in plain

    pytest.importorskip("mcp")
    server = build_server()

    async def call():
        return await server.call_tool("query_model", {
            "model": "gw_demo", "fact": "fact_orders",
            "measures": {"revenue": "sum"},
            "dimensions": ["dim_customer.region"],
            "order_by": ["-revenue"],
        })

    payload = anyio.run(call).model_dump(by_alias=True, exclude_none=True)
    sc = payload["structuredContent"]
    assert sc["fingerprint"] == out["fingerprint"]
    assert "tie-break" in sc["order_by_note"]
    assert sc["binding"]["query"]["order_by"] == sc["query"]["order_by"]


# ── Spec validation and rendering ──────────────────────────────────────────

def _spec(fact="fact_orders"):
    return {
        "name": "GW Spec",
        "sections": [{
            "type": "table",
            "title": "Revenue by region",
            "data": {
                "model": "gw_demo",
                "query": {
                    "fact": fact,
                    "measures": {"revenue": "sum"},
                    "dimensions": ["dim_customer.region"],
                },
            },
        }],
    }


# ── fetch_artifact: deliver the rendered bytes over MCP (WALL 2) ────────────


def test_fetch_is_path_guarded(gateway_model, tmp_path, monkeypatch):
    """Read-only and confined: no traversal, no absolute escape, no arbitrary
    file type, and nothing inside the installed package."""
    from tracebi.mcp_server import gateway_fetch_artifact

    monkeypatch.chdir(tmp_path)
    (tmp_path / "art.html").write_text("<html></html>", encoding="utf-8")
    (tmp_path / "secret.py").write_text("TOKEN = 'x'", encoding="utf-8")

    assert gateway_fetch_artifact("art.html")["ok"]              # in-root artifact OK
    assert not gateway_fetch_artifact("../escape.html")["ok"]    # traversal
    assert not gateway_fetch_artifact("/etc/hosts")["ok"]        # absolute escape
    assert not gateway_fetch_artifact("secret.py")["ok"]         # wrong extension
    assert not gateway_fetch_artifact("missing.html")["ok"]      # no such file

    import os

    import tracebi
    pkg = os.path.dirname(tracebi.__file__)
    inside_pkg = os.path.join(pkg, "web", "ui", "dist", "index.html")
    assert not gateway_fetch_artifact(inside_pkg)["ok"]          # installed package


# ── build_report: the publish step for the package lane ────────────────────

class TestBuildReport:
    def _package(self, tmp_path, monkeypatch, gateway_model):
        import tracebi.mcp_server as gw

        reports = tmp_path / "reports"
        pkg = reports / "gwpkg"
        pkg.mkdir(parents=True)
        (pkg / "report.json").write_text(json.dumps({
            "name": "gwpkg",
            "data": {"kpi": {"model": "gw_demo",
                             "query": {"fact": "fact_orders",
                                       "measures": ["total_revenue"]}}},
        }))
        (pkg / "template.html").write_text(
            "<html><head><title>g</title></head><body>"
            '<div data-tb-figure="value" data-tb-binding="kpi" '
            'data-tb-cell="total_revenue" id="fig-kpi"></div>'
            '<section data-tb-stage="exploration"><p>scratch</p></section>'
            "</body></html>"
        )
        monkeypatch.setenv("TRACEBI_REPORTS_DIR", str(reports))
        monkeypatch.setattr(gw, "_load_models",
                            lambda: {"gw_demo": gateway_model})
        return pkg

    def test_xlsx_writes_the_tables_and_fetch_returns_them(
            self, gateway_model, tmp_path, monkeypatch):
        """format=xlsx writes a workbook of the report's tables beside the
        HTML receipt. fetch_artifact returns that workbook base64-encoded.
        A path outside the output root is still refused."""
        import base64

        from openpyxl import load_workbook

        from tracebi.mcp_server import gateway_build_report, gateway_fetch_artifact

        self._package(tmp_path, monkeypatch, gateway_model)
        out_dir = tmp_path / "out"
        out = gateway_build_report(
            "gwpkg", output_dir=str(out_dir), format="xlsx")
        assert out["ok"], out.get("errors")
        xlsx = out_dir / "gwpkg.xlsx"
        assert out["xlsx_path"] == str(xlsx)
        assert xlsx.is_file()
        assert not (out_dir / "gwpkg.xlsx.manifest.json").exists()
        assert (out_dir / "gwpkg.html").is_file()
        assert (out_dir / "gwpkg.html.manifest.json").is_file()
        note = out["spreadsheet_note"]
        assert "no trace" in note and "not verifiable" in note
        assert "output_path" in note and "manifest_path" in note

        book = load_workbook(xlsx)
        cells = [
            c.value for row in book["Report"].iter_rows() for c in row
            if c.value is not None
        ]
        assert "kpi" in cells
        assert "total_revenue" in cells
        assert 900 in cells or 900.0 in cells

        monkeypatch.chdir(tmp_path)
        fetched = gateway_fetch_artifact(str(xlsx))
        assert fetched["ok"], fetched.get("errors")
        assert fetched["encoding"] == "base64"
        assert fetched["content_type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert base64.b64decode(fetched["content"]) == xlsx.read_bytes()
        assert fetched["bytes"] == xlsx.stat().st_size

        monkeypatch.setenv("TRACEBI_OUTPUT_ROOT", str(out_dir))
        escaped = gateway_build_report(
            "gwpkg", output_dir=str(tmp_path / "elsewhere"), format="xlsx")
        assert not escaped["ok"]
        assert "TRACEBI_OUTPUT_ROOT" in escaped["errors"][0]
        assert not (tmp_path / "elsewhere" / "gwpkg.xlsx").exists()
        outside = tmp_path / "elsewhere" / "gwpkg.xlsx"
        outside.parent.mkdir()
        outside.write_bytes(b"PK\x03\x04not-a-real-book")
        refused = gateway_fetch_artifact(str(outside))
        assert not refused["ok"]


    def test_refuses_writing_into_the_installed_package(self, gateway_model,
                                                        tmp_path, monkeypatch):
        """Always on: an agent must not write into the tracebi package (e.g.
        clobber web/ui/dist/index.html, which a server then serves)."""
        import pathlib

        import tracebi
        from tracebi.mcp_server import gateway_build_report

        self._package(tmp_path, monkeypatch, gateway_model)
        pkg_dist = str(
            pathlib.Path(tracebi.__file__).resolve().parent / "web" / "ui" / "dist")
        out = gateway_build_report("index", output_dir=pkg_dist)
        assert not out["ok"]
        assert "package" in out["errors"][0]

    def test_strict_confinement_refuses_output_outside_the_root(
            self, gateway_model, tmp_path, monkeypatch):
        """Opt-in: with TRACEBI_OUTPUT_ROOT set, an absolute path outside it or
        a traversal that escapes it is refused before any write."""
        from tracebi.mcp_server import gateway_build_report

        self._package(tmp_path, monkeypatch, gateway_model)
        monkeypatch.setenv("TRACEBI_OUTPUT_ROOT", str(tmp_path / "out"))
        for bad in ["/tmp/tracebi-escape-xyz", str(tmp_path / "elsewhere")]:
            out = gateway_build_report("gwpkg", output_dir=bad)
            assert not out["ok"], f"{bad!r} was not refused"
            assert "TRACEBI_OUTPUT_ROOT" in out["errors"][0]


# ── workbench_state: no report means the discovery session ─────────────────


def test_workbench_state_refuses_a_path_shaped_name(tmp_path, monkeypatch):
    """A caller-supplied report name must never become a path. Without the
    guard, report='/etc/x' or '../../x' escapes reports/ and collect_state
    reads — and executes report.py from — an attacker-chosen directory. The
    refusal must fire before any filesystem access."""
    import tracebi.mcp_server as gw
    from tracebi.mcp_server import gateway_workbench_state

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(gw, "_load_models", lambda: {})
    for payload in ["/etc/passwd", "../../etc/x", "../secrets", ".ssh/config",
                    "finance/../../x", "finance\\..\\x"]:
        out = gateway_workbench_state(report=payload)
        assert "errors" in out, f"{payload!r} was not refused: {out!r}"
        assert "invalid report name" in out["errors"][0]
    # A report in a folder is named by its path, and that is not refused: it
    # reaches the package lookup (and finds nothing here).
    out = gateway_workbench_state(report="finance/weekly")
    assert "invalid report name" not in " ".join(out.get("errors") or [])


# ── MCP 2.0 protocol features ────────────────────────────────────────────────
# The gateway advertises typed structured output, read-only tool annotations
# (the "read-and-compute only" refusal, in the protocol), reference resources,
# and an authoring prompt.


def test_a_schema_checking_client_accepts_every_tool(gateway_model, tmp_path,
                                                     monkeypatch):
    """Drive the real server through a schema-validating MCP client.

    The SDK fills every omitted result field with null and advertises that
    on the output schema. A client that checks structured content
    (``ClientSession.call_tool`` → ``validate_tool_result``) rejects a null
    the schema types as string, number, boolean, or array — and a success
    looks the same as a failure. Both the success path and the error
    envelope of every tool must come back as structured content.
    """
    pytest.importorskip("mcp")
    import anyio
    from mcp.client import Client

    from tracebi.mcp_server import build_server

    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("TRACEBI_MODELS_DIR", str(tmp_path / "models"))
    reports = tmp_path / "reports"
    pkg = reports / "gwpkg"
    pkg.mkdir(parents=True)
    (pkg / "report.json").write_text(json.dumps({
        "name": "gwpkg",
        "data": {"kpi": {"model": "gw_demo",
                         "query": {"fact": "fact_orders",
                                   "measures": ["total_revenue"]}}},
    }), encoding="utf-8")
    (pkg / "template.html").write_text(
        "<html><head><title>g</title></head><body>"
        '<div data-tb-figure="value" data-tb-binding="kpi" '
        'data-tb-cell="total_revenue" id="fig-kpi"></div>'
        "</body></html>",
        encoding="utf-8",
    )
    monkeypatch.setenv("TRACEBI_REPORTS_DIR", str(reports))
    out_dir = tmp_path / "built"
    server = build_server()

    async def drive():
        async with Client(server) as client:
            seen: list[str] = []

            async def call(name, arguments=None):
                result = await client.call_tool(name, arguments or {})
                seen.append(name)
                # A tool that raises is an isError result; the client does
                # not schema-check those. Every normal result — success or
                # the {ok, errors} envelope — must carry structured content.
                if not result.is_error:
                    assert result.structured_content is not None, name
                return result

            ctx = await call("get_context", {"brief": True})
            assert ctx.structured_content.get("tracebi_version")

            listed = await call("list_models")
            assert "gw_demo" in listed.structured_content["models"]

            described = await call("describe_model", {"model": "gw_demo"})
            assert described.structured_content.get("name") == "gw_demo"

            unknown = await call("describe_model", {"model": "no_such_model"})
            if unknown.is_error:
                text = unknown.content[0].text
                assert "Invalid structured content" not in text
                assert "not found" in text
            else:
                assert unknown.structured_content is not None

            tables = await call("describe_table")
            assert tables.structured_content.get("ok") is True
            one_table = await call("describe_table", {"table": "orders"})
            assert "ok" in one_table.structured_content

            query = await call("query_model", {
                "model": "gw_demo",
                "fact": "fact_orders",
                "measures": {"revenue": "sum"},
                "dimensions": ["dim_customer.region"],
            })
            assert query.structured_content.get("ok") is True
            assert query.structured_content.get("fingerprint")

            bad_query = await call("query_model", {
                "model": "gw_demo",
                "fact": "fact_orders",
                "measures": ["no_such_measure"],
            })
            assert bad_query.structured_content.get("ok") is False

            valid = await call("validate_report_spec", {"spec": _spec()})
            assert valid.structured_content.get("ok") is True
            invalid = await call("validate_report_spec",
                                 {"spec": _spec(fact="fact_nope")})
            assert invalid.structured_content.get("ok") is False

            rendered = await call("render_report_spec", {
                "spec": _spec(), "output_dir": str(out_dir),
            })
            assert rendered.structured_content.get("ok") is True
            refused_spec = await call("render_report_spec", {
                "spec": _spec(fact="fact_nope"), "output_dir": str(out_dir),
            })
            assert refused_spec.structured_content.get("ok") is False

            reports_listing = await call("list_reports")
            assert "reports" in reports_listing.structured_content

            built = await call("build_report", {
                "report": "gwpkg", "output_dir": str(out_dir),
            })
            assert built.structured_content.get("ok") is True, (
                built.structured_content)
            html_path = built.structured_content["output_path"]
            manifest_path = built.structured_content["manifest_path"]

            missing = await call("build_report", {
                "report": "missing", "output_dir": str(out_dir),
            })
            assert missing.structured_content.get("ok") is False

            fetched = await call("fetch_artifact", {"path": html_path})
            assert fetched.structured_content.get("ok") is True
            blocked = await call("fetch_artifact", {"path": "/etc/hosts"})
            assert blocked.structured_content.get("ok") is False

            verified = await call("verify_manifest", {"manifest": manifest_path})
            assert verified.structured_content.get("verdict")
            missing_manifest = await call("verify_manifest", {
                "manifest": str(tmp_path / "no-such.manifest.json"),
            })
            assert missing_manifest.structured_content.get("ok") is False

            discovery = await call("workbench_state")
            assert discovery.structured_content.get("mode") == "discovery"
            package_state = await call("workbench_state", {"report": "gwpkg"})
            assert package_state.structured_content.get("name")

            pin = await call("resolve_pin", {
                "report": "gwpkg", "pin_id": "no-such-pin",
            })
            assert pin.structured_content.get("ok") is False

            assert set(seen) == {
                "get_context", "list_models", "describe_model", "describe_table",
                "query_model", "validate_report_spec", "render_report_spec",
                "list_reports", "verify_manifest", "workbench_state",
                "resolve_pin", "build_report", "fetch_artifact",
            }

    anyio.run(drive)
